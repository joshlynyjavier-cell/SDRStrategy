"""Build the simplified Rippling new-business staffing model (formula-driven).

Usage: python3 -I analysis/build_model.py data/rippling_nb_dataset.csv model/Rippling_NB_Staffing_Model.xlsx
Colors: blue = editable input, black = formula, green = link to another tab.
"""
import sys
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

SRC, OUT = sys.argv[1], sys.argv[2]

FONT = "Arial"
F_IN = Font(name=FONT, size=10, color="0000FF")
F_FX = Font(name=FONT, size=10, color="000000")
F_LK = Font(name=FONT, size=10, color="008000")
F_NOTE = Font(name=FONT, size=9, italic=True, color="666666")
F_TITLE = Font(name=FONT, size=14, bold=True, color="1F2A44")
F_SEC = Font(name=FONT, size=11, bold=True, color="1F2A44")
F_H = Font(name=FONT, size=10, bold=True, color="FFFFFF")
FILL_H = PatternFill("solid", fgColor="1F2A44")
FILL_SEC = PatternFill("solid", fgColor="E8ECF4")
FILL_KEY = PatternFill("solid", fgColor="FFF2CC")
FILL_OUT = PatternFill("solid", fgColor="E2EFDA")

USD = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
PTS = '+0.0%;-0.0%;"-"'
N0 = '#,##0;(#,##0);"-"'
N1 = '#,##0.0;(#,##0.0);"-"'
N2 = '0.00;(0.00);"-"'

BUCKETS = ["1-10", "11-50", "51-100", "101-500", "501+"]
SEGS = ["Velocity (1-50)", "Emerging MM (51-100)", "Core MM (101-500)", "Enterprise (501+)"]
SEGMAP = {"1-10": (SEGS[0], "Velocity"), "11-50": (SEGS[0], "Velocity"), "51-100": (SEGS[1], "Mid-Market"),
          "101-500": (SEGS[2], "Mid-Market"), "501+": (SEGS[3], "Enterprise")}
SEGCOLS = ["C", "D", "E", "F"]  # segment columns on Inputs / Staffing

wb = Workbook()


def new_sheet(name, title, sub, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws["A1"], ws["A1"].font = title, F_TITLE
    ws["A2"], ws["A2"].font = sub, F_NOTE
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    return ws


def put(ws, ref, val, kind="fx", fmt=None, bold=False, fill=None):
    c = ws[ref]
    c.value = val
    base = {"in": F_IN, "fx": F_FX, "lk": F_LK, "note": F_NOTE}[kind]
    c.font = Font(name=FONT, size=base.size, color=base.color, bold=bold, italic=base.italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def header(ws, row, col, labels):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font, c.fill = F_H, FILL_H
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def section(ws, row, text, span):
    ws.cell(row=row, column=1, value=text).font = F_SEC
    for i in range(1, span + 1):
        ws.cell(row=row, column=i).fill = FILL_SEC


def note(ws, ref, text):
    put(ws, ref, text, "note")


# ------------------------------------------------------------------ README
ws = wb.active
ws.title = "README"
ws.sheet_view.showGridLines = False
ws.column_dimensions["B"].width = 22
ws.column_dimensions["C"].width = 115
ws["B1"], ws["B1"].font = "Rippling New-Business Staffing Model (simplified)", F_TITLE
ws["B2"], ws["B2"].font = "Built on the fictional case dataset: 12 months (Feb-2021 to Jan-2022) x 5 employee-size buckets.", F_NOTE
readme = [
    ("Purpose", "Analyze the new-business funnel by segment and size the SDR, AE and SC capacity each segment needs."),
    ("Colors", "Blue = editable input. Black = formula. Green = link from another tab. Yellow fill = key assumption."),
    ("Tabs", "Inputs -> Data -> Funnel -> Staffing -> Cost_Optional -> Checks."),
    ("Inputs", "Time windows, bucket-to-segment map, the 5 staffing assumptions, and the upside switch. ALL staffing assumptions are ILLUSTRATIVE, not Rippling data."),
    ("Data", "The 60 source rows from the CSV (blue) plus three formula columns: segment, early-period flag, baseline-period flag."),
    ("Funnel", "Funnel by segment and bucket, monthly trend, and early (Feb-Jul-21) vs baseline (Aug-Oct-21) comparison behind the three findings."),
    ("Staffing", "Baseline demand -> required FTE -> recommended headcount (ROUNDUP), ratios and expected ARR. Plus the upside case (lead-to-demo recovers)."),
    ("Cost_Optional", "Appendix only: a simple illustrative cost of the recommended headcount. It is not an ROI or payback analysis."),
    ("Checks", "Confirms the model reconciles to the source data. Every row should read TRUE."),
    ("Staffing formulas", "SDR FTE = leads / (SDR lead capacity x utilization). AE FTE = demos / (AE demo capacity x utilization). SC FTE = demos x SC engagement rate / (SC capacity x utilization)."),
    ("Rounding", "Dedicated roles: ROUNDUP per segment. Shared SC pool (Velocity + Emerging MM + Core MM): ROUNDUP of the combined FTE. Recommended headcount therefore always covers modeled workload."),
    ("Caveat 1", "Headcount shown is REQUIRED CAPACITY. Current headcount is not in the dataset, so this is not a hiring number. Enter current headcount on the Staffing tab to see a gap."),
    ("Caveat 2", "Funnel rates divide counts from the same month. They are not cohort conversion rates."),
    ("Caveat 3", "The lead-to-demo decline is a POTENTIAL top-of-funnel bottleneck. Possible causes include SDR capacity, speed-to-lead, lead quality and demo scheduling. The upside case is a sensitivity, not a forecast."),
    ("Caveat 4", "ARR per lead is revenue yield, not ROI. The dataset has no cost data."),
]
for i, (k, v) in enumerate(readme):
    r = 4 + i
    ws.cell(row=r, column=2, value=k).font = Font(name=FONT, size=10, bold=True)
    c = ws.cell(row=r, column=3, value=v)
    c.font, c.alignment = F_FX, Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 28 if len(v) > 115 else 15

# ------------------------------------------------------------------ Inputs
inp = new_sheet("Inputs", "Inputs (edit blue cells only)", "All staffing assumptions are ILLUSTRATIVE planning assumptions, not company data.",
                {"A": 3, "B": 52, "C": 17, "D": 19, "E": 17, "F": 17, "G": 90})
IN = {}
section(inp, 4, "1. Time windows", 7)
for r, key, lab, val, nt in [
    (5, "e_s", "Early period start", datetime(2021, 2, 1), "Early period = Feb-Jul 2021 (first 6 months of data)."),
    (6, "e_e", "Early period end", datetime(2021, 7, 1), ""),
    (7, "b_s", "Baseline period start", datetime(2021, 8, 1), "Baseline = Aug-Oct 2021, the latest 3 months before the Nov-Dec dip."),
    (8, "b_e", "Baseline period end", datetime(2021, 10, 1), "Nov-Dec excluded: seen once only, so it is not confirmed seasonality. Jan-22 rebounded to a record."),
]:
    put(inp, f"B{r}", lab)
    put(inp, f"C{r}", val, "in", "mmm-yy", fill=FILL_KEY)
    note(inp, f"G{r}", nt)
    IN[key] = f"Inputs!$C${r}"

section(inp, 10, "2. Segment map (dataset bucket -> segment)", 7)
header(inp, 11, 2, ["Employee bucket", "Segment", "Sales motion"])
for i, b in enumerate(BUCKETS):
    r = 12 + i
    put(inp, f"B{r}", b, "in")
    put(inp, f"C{r}", SEGMAP[b][0], "in")
    put(inp, f"D{r}", SEGMAP[b][1], "in")
note(inp, "G12", "Boundaries sit where the data changes: ARR/logo jumps ~2.3x at 50, conversion halves at 100 and again at 500.")

section(inp, 18, "3. Staffing assumptions (ILLUSTRATIVE; per ramped employee per month)", 7)
header(inp, 19, 2, ["Assumption"] + SEGS + ["Why this value (illustrative)"])
ASSUME = [
    ("sdr_cap", "SDR lead capacity (high-intent leads worked / month)", [300, 200, 175, 125], N0,
     "Larger accounts need more research and follow-up per lead. 200/month = ~10 new leads per working day."),
    ("ae_cap", "AE demo capacity (demos held / month)", [45, 30, 20, 12], N0,
     "Includes follow-up on the opps each demo creates. Velocity ~2 demos/day; Enterprise demos are multi-session."),
    ("sc_rate", "SC engagement rate (% of demos needing an SC)", [0.05, 0.15, 0.50, 1.00], PCT,
     "Rises with complexity seen in the data: demo->opp falls from ~80% to ~65% and demos per logo rise from 3.7 to 7.8."),
    ("sc_cap", "SC engagement capacity (SC-supported demos / month)", [30, 30, 30, 20], N0,
     "Enterprise engagements are deeper (discovery, custom demo, security review)."),
    ("util", "Target utilization", [0.85, 0.85, 0.85, 0.85], PCT,
     "15% buffer for PTO, training, admin and demand spikes."),
]
for i, (key, lab, vals, fmt, why) in enumerate(ASSUME):
    r = 20 + i
    put(inp, f"B{r}", lab)
    for col, v in zip(SEGCOLS, vals):
        put(inp, f"{col}{r}", v, "in", fmt, fill=FILL_KEY)
    note(inp, f"G{r}", why)
    IN[key] = r

section(inp, 26, "4. Upside sensitivity", 7)
put(inp, "B27", "Lead-to-demo recovers to early-period level? (1 = yes)")
for col, v in zip(SEGCOLS, [1, 1, 1, 0]):
    put(inp, f"{col}27", v, "in", N0, fill=FILL_KEY)
note(inp, "G27", "Enterprise = 0: its lead-to-demo did not decline. The upside is a 'what if', not a forecast.")
IN["up"] = 27

section(inp, 29, "5. Shared resource pool", 7)
note(inp, "B30", "One shared SC pool covers Velocity, Emerging MM and Core MM. Below 100 employees each segment needs less than one SC, so the pool is sized on")
note(inp, "B31", "the combined FTE (ROUNDUP). Enterprise SCs are dedicated. All SDR and AE roles are dedicated per segment.")
note(inp, "B32", "In practice, Velocity and Emerging MM SDRs could share one inbound queue, with 51-100 leads routed first.")

seg_rng = "Inputs!$B$12:$B$16"

# ------------------------------------------------------------------ Data
raw = pd.read_csv(SRC)
raw.columns = ["month", "fte", "leads", "demos", "qo", "logos", "arr"]
raw["arr"] = raw["arr"].str.replace(r"[$,]", "", regex=True).astype(float)
raw["month"] = pd.to_datetime(raw["month"])
dt = wb.create_sheet("Data")
dt.sheet_view.showGridLines = False
header(dt, 1, 1, ["Month", "Employee bucket", "Leads", "Demos", "Qualified opps", "Logos", "ARR ($)", "Segment", "Early (1/0)", "Baseline (1/0)"])
for c, w in zip("ABCDEFGHIJ", [10, 15, 9, 9, 13, 9, 12, 22, 11, 13]):
    dt.column_dimensions[c].width = w
for i, row in raw.iterrows():
    r = i + 2
    put(dt, f"A{r}", row.month.to_pydatetime(), "in", "mmm-yy")
    put(dt, f"B{r}", row.fte, "in")
    for col, k in zip("CDEF", ["leads", "demos", "qo", "logos"]):
        put(dt, f"{col}{r}", int(row[k]), "in", N0)
    put(dt, f"G{r}", float(row.arr), "in", USD)
    put(dt, f"H{r}", f"=INDEX(Inputs!$C$12:$C$16,MATCH(B{r},{seg_rng},0))", "lk")
    put(dt, f"I{r}", f"=IF(AND(A{r}>={IN['e_s']},A{r}<={IN['e_e']}),1,0)")
    put(dt, f"J{r}", f"=IF(AND(A{r}>={IN['b_s']},A{r}<={IN['b_e']}),1,0)")
dt.freeze_panes = "A2"
NR = len(raw) + 1
R = lambda col: f"Data!${col}$2:${col}${NR}"
COL = {"leads": "C", "demos": "D", "qo": "E", "logos": "F", "arr": "G"}

# ------------------------------------------------------------------ Funnel
fn = new_sheet("Funnel", "Funnel analysis (all values are formulas on the Data tab)",
               "Rates = summed numerator / summed denominator. Same-month ratios, not cohort conversion.",
               {"A": 30, **{c: 12 for c in "BCDEFGHIJKLMNO"}})
FCOLS = ["Leads", "Demos", "Qual opps", "Logos", "ARR", "Lead->Demo", "Demo->Opp", "Opp->Win", "Lead->Win", "ARR / logo", "ARR / lead",
         "Lead share", "ARR share", "Demos / logo"]


def funnel_table(top, title, keys, crit_col):
    section(fn, top, title, 15)
    header(fn, top + 1, 1, ["Segment" if crit_col == "H" else "Bucket"] + FCOLS)
    first, last = top + 2, top + 1 + len(keys)
    tot = last + 1
    for i, k in enumerate(keys + ["Total"]):
        r = top + 2 + i
        put(fn, f"A{r}", k, bold=(k == "Total"))
        for j, m in enumerate(["leads", "demos", "qo", "logos", "arr"]):
            c = "BCDEF"[j]
            f = f"=SUM({c}{first}:{c}{last})" if k == "Total" else f"=SUMIFS({R(COL[m])},{R(crit_col)},$A{r})"
            put(fn, f"{c}{r}", f, fmt=USD if m == "arr" else N0, bold=(k == "Total"))
        for c, f, fmt in [("G", f"=C{r}/B{r}", PCT), ("H", f"=D{r}/C{r}", PCT), ("I", f"=E{r}/D{r}", PCT), ("J", f"=E{r}/B{r}", PCT),
                          ("K", f"=F{r}/E{r}", USD), ("L", f"=F{r}/B{r}", USD), ("M", f"=B{r}/B${tot}", PCT), ("N", f"=F{r}/F${tot}", PCT),
                          ("O", f"=C{r}/E{r}", N1)]:
            put(fn, f"{c}{r}", f, fmt=fmt, bold=(k == "Total"))
    return first, tot


SEG_F0, SEG_FT = funnel_table(4, "A. 12-month funnel by segment (Feb-21 to Jan-22)", SEGS, "H")
BKT_F0, BKT_FT = funnel_table(12, "B. 12-month funnel by employee bucket (detail)", BUCKETS, "B")

section(fn, 21, "C. Monthly trend (all segments)", 15)
header(fn, 22, 1, ["Month", "Leads", "Demos", "Logos", "ARR", "Lead->Demo", "Demo->Win", "ARR / logo", "Leads from 101+ (share)"])
months = sorted(raw["month"].unique())
for i, m in enumerate(months):
    r = 23 + i
    put(fn, f"A{r}", pd.Timestamp(m).to_pydatetime(), "in", "mmm-yy")
    for c, k in zip("BCDE", ["leads", "demos", "logos", "arr"]):
        put(fn, f"{c}{r}", f"=SUMIFS({R(COL[k])},{R('A')},$A{r})", fmt=USD if k == "arr" else N0)
    put(fn, f"F{r}", f"=C{r}/B{r}", fmt=PCT)
    put(fn, f"G{r}", f"=D{r}/C{r}", fmt=PCT)
    put(fn, f"H{r}", f"=E{r}/D{r}", fmt=USD)
    put(fn, f"I{r}", f"=(SUMIFS({R('C')},{R('A')},$A{r},{R('B')},\"101-500\")+SUMIFS({R('C')},{R('A')},$A{r},{R('B')},\"501+\"))/B{r}", fmt=PCT)
put(fn, "A35", "Total", bold=True)
for c in "BCDE":
    put(fn, f"{c}35", f"=SUM({c}23:{c}34)", fmt=USD if c == "E" else N0, bold=True)
note(fn, "A36", "Nov-Dec 2021: leads fell ~47% vs Oct and every conversion rate dipped; Jan-22 set records. With one year of data, this is not confirmed seasonality.")

section(fn, 38, "D. Early period (Feb-Jul-21) vs baseline (Aug-Oct-21): average month", 15)
header(fn, 39, 1, ["Metric"] + SEGS + ["Total", "Reading"])
put(fn, "A40", "Segment")
for col, s in zip(["B", "C", "D", "E"], SEGS):
    put(fn, f"{col}40", s, "lk")
put(fn, "F40", "Total", bold=True)


def avg(metric, flag, seg_col):
    crit = f",{R('H')},{seg_col}$40" if seg_col else ""
    return f"(SUMIFS({R(COL[metric])},{R(flag)},1{crit})/(COUNTIFS({R(flag)},1)/5))"


rows = [("el", "Leads / month (early)", N0), ("bl", "Leads / month (baseline)", N0), ("gl", "  Lead growth", PCT),
        ("ea", "ARR / month (early)", USD), ("ba", "ARR / month (baseline)", USD), ("ga", "  ARR growth", PCT),
        ("ep", "ARR / logo (early)", USD), ("bp", "ARR / logo (baseline)", USD), ("gp", "  ARR / logo growth", PCT),
        ("ed", "Lead->Demo (early)", PCT), ("bd", "Lead->Demo (baseline)", PCT), ("cd", "  Change (pts)", PTS),
        ("eq", "Demo->Opp (early)", PCT), ("bq", "Demo->Opp (baseline)", PCT), ("cq", "  Change (pts)", PTS),
        ("ew", "Opp->Win (early)", PCT), ("bw", "Opp->Win (baseline)", PCT), ("cw", "  Change (pts)", PTS)]
EV = {key: 41 + i for i, (key, _, _) in enumerate(rows)}
PAIR = {"l": ("el", "bl"), "a": ("ea", "ba"), "p": ("ep", "bp"), "d": ("ed", "bd"), "q": ("eq", "bq"), "w": ("ew", "bw")}
RATIO = {"p": ("arr", "logos"), "d": ("demos", "leads"), "q": ("qo", "demos"), "w": ("logos", "qo")}
for key, lab, fmt in rows:
    r = EV[key]
    put(fn, f"A{r}", lab)
    for col in ["B", "C", "D", "E", "F"]:
        sc = None if col == "F" else col
        kind, m = key[0], key[1]
        if kind in ("e", "b"):
            flag = "I" if kind == "e" else "J"
            if m == "l":
                f = "=" + avg("leads", flag, sc)
            elif m == "a":
                f = "=" + avg("arr", flag, sc)
            else:
                n, d = RATIO[m]
                f = f"={avg(n, flag, sc)}/{avg(d, flag, sc)}"
        else:
            e_key, b_key = PAIR[m]
            f = f"={col}{EV[b_key]}/{col}{EV[e_key]}-1" if kind == "g" else f"={col}{EV[b_key]}-{col}{EV[e_key]}"
        put(fn, f"{col}{r}", f, fmt=fmt, bold=(col == "F"))
note(fn, f"G{EV['gl']}", "Demand grew in every segment except Velocity; Enterprise grew fastest.")
note(fn, f"G{EV['cd']}", "Lead->Demo fell in the three segments under 500 employees. Enterprise held.")
note(fn, f"G{EV['cq']}", "Post-demo conversion broadly held, so the slippage is concentrated at the top of the funnel.")
note(fn, f"G{EV['cw']}", "Win rates broadly stable. Larger segments convert at roughly half the rate of smaller ones.")
r = EV["cw"] + 1
put(fn, f"A{r}", "Lead share from 101+ employees")
put(fn, f"B{r}", f"=({avg('leads', 'I', 'D')}+{avg('leads', 'I', 'E')})/{avg('leads', 'I', None)}", fmt=PCT)
put(fn, f"C{r}", f"=({avg('leads', 'J', 'D')}+{avg('leads', 'J', 'E')})/{avg('leads', 'J', None)}", fmt=PCT)
note(fn, f"D{r}", "B = early period, C = baseline. The lead mix is shifting toward larger, lower-converting accounts.")
EV["mix"] = r

ch = LineChart()
ch.title, ch.height, ch.width = "Monthly leads vs demos", 7.5, 15
ch.add_data(Reference(fn, min_col=2, max_col=3, min_row=22, max_row=34), titles_from_data=True)
ch.set_categories(Reference(fn, min_col=1, min_row=23, max_row=34))
fn.add_chart(ch, "Q4")

# ------------------------------------------------------------------ Staffing
st = new_sheet("Staffing", "Staffing model: required capacity by segment",
               "Required capacity, NOT net-new hires (current headcount is not in the dataset). The allocation follows from the stated illustrative assumptions; it is not a claim of optimality.",
               {"A": 3, "B": 52, "C": 17, "D": 19, "E": 17, "F": 17, "G": 13, "H": 72})
header(st, 4, 2, ["Line item"] + SEGS + ["Total", "Formula / note"])


def srow(r, lab, fseg, fmt, total="sum", nt="", kind="fx", fill=None):
    put(st, f"B{r}", lab)
    for col in SEGCOLS:
        put(st, f"{col}{r}", fseg(col), kind, fmt, fill=fill)
    if total == "sum":
        put(st, f"G{r}", f"=SUM(C{r}:F{r})", fmt=fmt, bold=True, fill=fill)
    elif total:
        put(st, f"G{r}", total, fmt=fmt, bold=True, fill=fill)
    if nt:
        note(st, f"H{r}", nt)


IA = lambda key, col: f"Inputs!{col}${IN[key]}"
bavg = lambda m, col: f"(SUMIFS({R(COL[m])},{R('H')},{col}$5,{R('J')},1)/(COUNTIFS({R('J')},1)/5))"
eavg = lambda m, col: f"(SUMIFS({R(COL[m])},{R('H')},{col}$5,{R('I')},1)/(COUNTIFS({R('I')},1)/5))"

put(st, "B5", "Segment")
for col, s in zip(SEGCOLS, SEGS):
    put(st, f"{col}5", s, "lk", bold=True)

section(st, 6, "1. Baseline demand: average month, Aug-Oct 2021 (from Data)", 8)
srow(7, "Leads / month", lambda c: "=" + bavg("leads", c), N1, nt="Average of the baseline months")
srow(8, "Demos / month", lambda c: "=" + bavg("demos", c), N1)
srow(9, "Lead->Demo", lambda c: f"={c}8/{c}7", PCT, "=G8/G7", nt="Demos / leads")
srow(10, "Demo->Win (logos / demos)", lambda c: f"=({bavg('logos', c)})/{c}8", PCT, "=SUMPRODUCT(C8:F8,C10:F10)/G8", nt="Logos / demos")
srow(11, "ARR / logo", lambda c: f"=({bavg('arr', c)})/({bavg('logos', c)})", USD,
     "=SUMPRODUCT(C8:F8,C10:F10,C11:F11)/SUMPRODUCT(C8:F8,C10:F10)", nt="ARR / logos")

section(st, 13, "2. Staffing assumptions (linked from Inputs; illustrative)", 8)
srow(14, "SDR lead capacity / month", lambda c: f"={IA('sdr_cap', c)}", N0, None, kind="lk")
srow(15, "AE demo capacity / month", lambda c: f"={IA('ae_cap', c)}", N0, None, kind="lk")
srow(16, "SC engagement rate", lambda c: f"={IA('sc_rate', c)}", PCT, None, kind="lk")
srow(17, "SC engagement capacity / month", lambda c: f"={IA('sc_cap', c)}", N0, None, kind="lk")
srow(18, "Target utilization", lambda c: f"={IA('util', c)}", PCT, None, kind="lk")

section(st, 20, "3. Baseline required FTE (fractional workload)", 8)
srow(21, "SDR FTE", lambda c: f"={c}7/({c}14*{c}18)", N2, nt="Leads / (SDR capacity x utilization)")
srow(22, "AE FTE", lambda c: f"={c}8/({c}15*{c}18)", N2, nt="Demos / (AE capacity x utilization)")
srow(23, "SC FTE", lambda c: f"={c}8*{c}16/({c}17*{c}18)", N2, nt="Demos x SC rate / (SC capacity x utilization)")

section(st, 25, "4. Recommended headcount, baseline (ROUNDUP; covers modeled workload)", 8)
srow(26, "SDR (dedicated per segment)", lambda c: f"=ROUNDUP({c}21,0)", N0, nt="ROUNDUP of each segment's SDR FTE", fill=FILL_OUT)
srow(27, "AE (dedicated per segment)", lambda c: f"=ROUNDUP({c}22,0)", N0, nt="ROUNDUP of each segment's AE FTE", fill=FILL_OUT)
put(st, "B28", "SC: shared pool (Velocity + Emerging MM + Core MM) | Enterprise")
put(st, "C28", "=ROUNDUP(SUM(C23:E23),0)", fmt=N0, fill=FILL_OUT)
st.merge_cells("C28:E28")
put(st, "F28", "=ROUNDUP(F23,0)", fmt=N0, fill=FILL_OUT)
put(st, "G28", "=C28+F28", fmt=N0, bold=True, fill=FILL_OUT)
note(st, "H28", "Pool = ROUNDUP of combined FTE in C23:E23. Enterprise = ROUNDUP of its own FTE.")
put(st, "B29", "Total recommended headcount", bold=True)
put(st, "G29", "=G26+G27+G28", fmt=N0, bold=True, fill=FILL_OUT)
srow(30, "AEs per SDR", lambda c: f"={c}27/{c}26", N1, "=G27/G26", nt="Coverage ratio")
put(st, "B31", "AEs per SC")
put(st, "C31", "=SUM(C27:E27)/C28", fmt=N1)
st.merge_cells("C31:E31")
put(st, "F31", "=F27/F28", fmt=N1)
put(st, "G31", "=G27/G28", fmt=N1, bold=True)
note(st, "H31", "Shared pool: commercial AEs per pooled SC")
srow(32, "AE utilization at recommended headcount", lambda c: f"={c}8/({c}27*{c}15)", PCT, "=G8/SUMPRODUCT(C27:F27,C15:F15)",
     nt="Demos / (AE headcount x capacity). Below 85% = headroom created by rounding up.")

section(st, 34, "5. Expected ARR, baseline (observed conversion; no improvement assumed)", 8)
srow(35, "New ARR / month", lambda c: f"={c}8*{c}10*{c}11", USD, nt="Demos x demo->win x ARR/logo (= observed baseline ARR)")
srow(36, "Annualized new ARR (x12)", lambda c: f"={c}35*12", USD, nt="Run-rate. If the Nov-Dec dip recurs, annual ARR would be lower.")
srow(37, "ARR per AE (recommended headcount)", lambda c: f"={c}36/{c}27", USD, "=G36/G27", nt="Productivity check only")

section(st, 39, "6. Upside sensitivity: lead-to-demo recovers to the early-period level (NOT a forecast)", 8)
srow(40, "Early-period Lead->Demo (Feb-Jul-21)", lambda c: f"=({eavg('demos', c)})/({eavg('leads', c)})", PCT, None)
srow(41, "Recovery applied? (from Inputs)", lambda c: f"={IA('up', c)}", N0, None, kind="lk")
srow(42, "Upside Lead->Demo", lambda c: f"=IF({c}41=1,MAX({c}9,{c}40),{c}9)", PCT, "=G43/G7")
srow(43, "Upside demos / month", lambda c: f"={c}7*{c}42", N1, nt="Same leads as baseline; a higher share reaches a demo")
srow(44, "Upside AE FTE", lambda c: f"={c}43/({c}15*{c}18)", N2)
srow(45, "Upside SC FTE", lambda c: f"={c}43*{c}16/({c}17*{c}18)", N2)
srow(46, "Upside SDR headcount", lambda c: f"={c}26", N0, nt="Unchanged: SDRs are sized to leads, which do not change", fill=FILL_OUT)
srow(47, "Upside AE headcount", lambda c: f"=ROUNDUP({c}44,0)", N0, fill=FILL_OUT)
put(st, "B48", "Upside SC headcount (pool | Enterprise)")
put(st, "C48", "=ROUNDUP(SUM(C45:E45),0)", fmt=N0, fill=FILL_OUT)
st.merge_cells("C48:E48")
put(st, "F48", "=ROUNDUP(F45,0)", fmt=N0, fill=FILL_OUT)
put(st, "G48", "=C48+F48", fmt=N0, bold=True, fill=FILL_OUT)
put(st, "B49", "Upside total headcount", bold=True)
put(st, "G49", "=G46+G47+G48", fmt=N0, bold=True, fill=FILL_OUT)
srow(50, "Upside annualized new ARR", lambda c: f"={c}43*{c}10*{c}11*12", USD, nt="Upside demos x baseline demo->win x ARR/logo x 12")
srow(51, "Change in ARR vs baseline", lambda c: f"={c}50-{c}36", USD, nt="Only if the decline is fixable top-of-funnel friction. If it reflects lead quality, there is no uplift.")
srow(52, "Change in AE headcount vs baseline", lambda c: f"={c}47-{c}27", N0)
put(st, "B53", "Change in total headcount vs baseline", bold=True)
put(st, "G53", "=G49-G29", fmt=N0, bold=True)

section(st, 55, "7. Optional: enter current ramped headcount to see the hiring gap (blank = unknown)", 8)
for j, role in enumerate(["SDR", "AE"]):
    r = 56 + j
    put(st, f"B{r}", f"Current {role} (enter)")
    for col in SEGCOLS:
        put(st, f"{col}{r}", None, "in", N0, fill=FILL_KEY)
    g = 59 + j
    put(st, f"B{g}", f"{role} gap = recommended - current")
    for col in SEGCOLS:
        put(st, f"{col}{g}", f'=IF({col}{r}="","n/a",{col}{26 + j}-{col}{r})', fmt=N0)
put(st, "B58", "Current SC (enter: shared pool in C, Enterprise in F)")
put(st, "C58", None, "in", N0, fill=FILL_KEY)
put(st, "F58", None, "in", N0, fill=FILL_KEY)
put(st, "B61", "SC gap = recommended - current")
put(st, "C61", '=IF(C58="","n/a",C28-C58)', fmt=N0)
put(st, "F61", '=IF(F58="","n/a",F28-F58)', fmt=N0)
note(st, "B62", "Hiring need = required capacity - current ramped capacity (+ attrition backfill). Not knowable from this dataset.")

section(st, 64, "8. How staffing ties to the findings and priorities", 8)
links = [
    ("Velocity (1-50)", "Efficient but low value ($3.8K ARR/logo); the 1-10 bucket is shrinking. Lean AE coverage, automated scheduling, minimal SC."),
    ("Emerging MM (51-100)", "Strongest observed revenue yield per lead, and the largest lead-to-demo slippage. Dedicated AEs plus SDR speed-to-lead coverage."),
    ("Core MM (101-500)", "Largest ARR pool and growing. Lower conversion and more demos per logo, so it gets the largest AE team and SC support on half of demos."),
    ("Enterprise (501+)", "Fastest growing but lowest conversion. Full SC support and dedicated SDRs; hold expansion until win-rate improvement is shown."),
]
for i, (s_, t) in enumerate(links):
    put(st, f"B{65 + i}", s_, bold=True)
    c = st[f"C{65 + i}"]
    c.value, c.font = t, F_FX
    st.merge_cells(f"C{65 + i}:H{65 + i}")
st.freeze_panes = "C6"

# ------------------------------------------------------------------ Cost_Optional
co = new_sheet("Cost_Optional", "OPTIONAL appendix: illustrative staffing cost of recommended headcount",
               "Illustrative SF fully loaded annual cost (salary + variable at target + benefits/taxes). Validate with Comp. Not an ROI, CAC or payback analysis.",
               {"A": 3, "B": 34, "C": 18, "D": 18, "E": 18, "F": 18, "G": 18})
header(co, 4, 2, ["Role", "Loaded cost / FTE (illustrative)", "Baseline headcount", "Baseline cost", "Upside headcount", "Upside cost"])
for i, (role, cost, hb, hu) in enumerate([("SDR", 115000, "G26", "G46"), ("AE (blended across segments)", 250000, "G27", "G47"), ("SC", 265000, "G28", "G48")]):
    r = 5 + i
    put(co, f"B{r}", role)
    put(co, f"C{r}", cost, "in", USD, fill=FILL_KEY)
    put(co, f"D{r}", f"=Staffing!{hb}", "lk", N0)
    put(co, f"E{r}", f"=C{r}*D{r}", fmt=USD)
    put(co, f"F{r}", f"=Staffing!{hu}", "lk", N0)
    put(co, f"G{r}", f"=C{r}*F{r}", fmt=USD)
put(co, "B8", "Total", bold=True)
for c in "DEFG":
    put(co, f"{c}8", f"=SUM({c}5:{c}7)", fmt=USD if c in "EG" else N0, bold=True)
note(co, "B10", "Use: a sense of scale for the required capacity. Current headcount is unknown, so this is NOT incremental spend.")
note(co, "B11", "AE pay varies by segment in practice (Enterprise higher, Velocity lower). One blended figure keeps this optional view simple.")

# ------------------------------------------------------------------ Checks
ck = new_sheet("Checks", "Reconciliation checks", "Every row should read TRUE.", {"A": 3, "B": 72, "C": 16, "D": 16, "E": 10})
header(ck, 4, 2, ["Check", "Value A", "Value B", "Pass?"])
checks = [
    ("Data has 60 rows (12 months x 5 buckets)", f"=COUNTA({R('B')})", "=60", "eq"),
    ("Funnel segment total leads = sum of Data leads", f"=Funnel!B{SEG_FT}", f"=SUM({R('C')})", "eq"),
    ("Funnel segment total ARR = sum of Data ARR", f"=Funnel!F{SEG_FT}", f"=SUM({R('G')})", "eq"),
    ("Funnel bucket total ARR = segment total ARR", f"=Funnel!F{BKT_FT}", f"=Funnel!F{SEG_FT}", "eq"),
    ("Monthly trend ARR = segment total ARR", "=Funnel!E35", f"=Funnel!F{SEG_FT}", "eq"),
    ("Staffing baseline leads/month = Data baseline leads / months", "=Staffing!G7", f"=SUMIFS({R('C')},{R('J')},1)/(COUNTIFS({R('J')},1)/5)", "eq"),
    ("Staffing baseline ARR/month = Data baseline ARR / months", "=Staffing!G35", f"=SUMIFS({R('G')},{R('J')},1)/(COUNTIFS({R('J')},1)/5)", "eq"),
    ("Staffing baseline leads = Funnel baseline leads", "=Staffing!G7", f"=Funnel!F{EV['bl']}", "eq"),
    ("SDR headcount covers SDR FTE (every segment)", "=SUMPRODUCT((Staffing!C26:F26<Staffing!C21:F21)*1)", "=0", "eq"),
    ("AE headcount covers AE FTE (every segment)", "=SUMPRODUCT((Staffing!C27:F27<Staffing!C22:F22)*1)", "=0", "eq"),
    ("SC pool covers pooled SC FTE", "=Staffing!C28", "=SUM(Staffing!C23:E23)", "ge"),
    ("Enterprise SC covers Enterprise SC FTE", "=Staffing!F28", "=Staffing!F23", "ge"),
    ("Total headcount = SDR + AE + SC", "=Staffing!G29", "=SUM(Staffing!G26:G28)", "eq"),
    ("Upside headcount >= baseline headcount", "=Staffing!G49", "=Staffing!G29", "ge"),
    ("Upside ARR >= baseline ARR", "=Staffing!G50", "=Staffing!G36", "ge"),
    ("Cost tab headcount = Staffing headcount", "=Cost_Optional!D8", "=Staffing!G29", "eq"),
]
for i, (lab, a, b, op) in enumerate(checks):
    r = 5 + i
    put(ck, f"B{r}", lab)
    put(ck, f"C{r}", a, "lk", N2)
    put(ck, f"D{r}", b, fmt=N2)
    put(ck, f"E{r}", f"=ABS(C{r}-D{r})<0.01" if op == "eq" else f"=C{r}>=D{r}-0.0001")
last = 5 + len(checks) - 1
put(ck, f"B{last + 2}", "ALL CHECKS PASS", bold=True)
put(ck, f"E{last + 2}", f"=AND(E5:E{last})", bold=True)

order = ["README", "Inputs", "Data", "Funnel", "Staffing", "Cost_Optional", "Checks"]
wb._sheets = [wb[n] for n in order]
for w in wb.worksheets:
    w.sheet_properties.tabColor = {"Inputs": "0000FF", "Staffing": "008000", "Checks": "C00000"}.get(w.title, "1F2A44")
wb.save(OUT)
print("saved", OUT, "| funnel rows", SEG_F0, SEG_FT, EV)
